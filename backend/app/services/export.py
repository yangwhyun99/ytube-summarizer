"""PDF / DOCX 내보내기 서비스"""

import io
from dataclasses import dataclass


@dataclass
class ExportSection:
    title: str
    content: str
    timestamp_start: float | None = None
    timestamp_end: float | None = None


def _format_time(seconds: float | None) -> str:
    if seconds is None:
        return ""
    m = int(seconds // 60)
    s = int(seconds % 60)
    return f"{m}:{s:02d}"


def export_pdf(title: str, sections: list[ExportSection]) -> bytes:
    """요약을 PDF로 내보내기"""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.enums import TA_LEFT
    import os

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=20 * mm, bottomMargin=20 * mm)
    styles = getSampleStyleSheet()

    # 한국어 폰트 등록 시도
    font_name = "Helvetica"
    korean_font_paths = [
        "C:/Windows/Fonts/malgun.ttf",  # Windows 맑은 고딕
        "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",  # Linux
        "/System/Library/Fonts/AppleGothic.ttf",  # Mac
    ]
    for fp in korean_font_paths:
        if os.path.exists(fp):
            try:
                pdfmetrics.registerFont(TTFont("KoreanFont", fp))
                font_name = "KoreanFont"
                break
            except Exception:
                continue

    title_style = ParagraphStyle(
        "CustomTitle", parent=styles["Title"], fontName=font_name, fontSize=18
    )
    heading_style = ParagraphStyle(
        "CustomHeading", parent=styles["Heading2"], fontName=font_name, fontSize=14
    )
    body_style = ParagraphStyle(
        "CustomBody", parent=styles["Normal"], fontName=font_name, fontSize=11,
        leading=16, alignment=TA_LEFT,
    )
    time_style = ParagraphStyle(
        "TimeStyle", parent=styles["Normal"], fontName=font_name, fontSize=9,
        textColor="grey",
    )

    elements = []
    elements.append(Paragraph(title, title_style))
    elements.append(Spacer(1, 10 * mm))

    for section in sections:
        elements.append(Paragraph(section.title, heading_style))

        if section.timestamp_start is not None:
            ts = _format_time(section.timestamp_start)
            te = _format_time(section.timestamp_end)
            time_text = f"[{ts}" + (f" - {te}]" if te else "]")
            elements.append(Paragraph(time_text, time_style))

        # 줄바꿈 처리
        for line in section.content.split("\n"):
            if line.strip():
                elements.append(Paragraph(line, body_style))
        elements.append(Spacer(1, 5 * mm))

    doc.build(elements)
    return buf.getvalue()


def export_docx(title: str, sections: list[ExportSection]) -> bytes:
    """요약을 DOCX로 내보내기"""
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()

    # 제목
    doc.add_heading(title, level=0)

    for section in sections:
        # 섹션 제목
        doc.add_heading(section.title, level=2)

        # 타임스탬프
        if section.timestamp_start is not None:
            ts = _format_time(section.timestamp_start)
            te = _format_time(section.timestamp_end)
            time_text = f"[{ts}" + (f" - {te}]" if te else "]")
            p = doc.add_paragraph()
            run = p.add_run(time_text)
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(128, 128, 128)

        # 내용
        doc.add_paragraph(section.content)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
